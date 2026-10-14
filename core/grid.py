"""Read a spreadsheet file into a plain grid of cell values plus its merged ranges.

Every later step (layout detection, parsing) works on this grid, so the three
file formats (.xls, .xlsx, .csv) are handled in one place.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from dataclasses import dataclass, field
from pathlib import Path

# (first_row, last_row_exclusive, first_col, last_col_exclusive), as xlrd reports them
MergedRange = tuple[int, int, int, int]

SUPPORTED_SUFFIXES = (".xls", ".xlsx", ".csv")


class UnsupportedFile(ValueError):
    """The file type is not one the app can read."""


@dataclass
class Grid:
    """Cell values of the first sheet. Values are str, int, float, datetime or None."""

    cells: list[list[object]]
    merged: list[MergedRange] = field(default_factory=list)
    source_format: str = ""

    @property
    def nrows(self) -> int:
        return len(self.cells)

    @property
    def ncols(self) -> int:
        return max((len(row) for row in self.cells), default=0)

    def value(self, r: int, c: int) -> object:
        if r >= self.nrows or c >= len(self.cells[r]):
            return None
        return self.cells[r][c]

    def text(self, r: int, c: int) -> str:
        """The cell as trimmed text; integral floats lose their '.0'."""
        return cell_text(self.value(r, c))

    def merged_origin(self, r: int, c: int) -> tuple[int, int]:
        """The top-left cell of the merged range covering (r, c), or (r, c) itself."""
        for r0, r1, c0, c1 in self.merged:
            if r0 <= r < r1 and c0 <= c < c1:
                return r0, c0
        return r, c

    def row_is_blank(self, r: int) -> bool:
        return all(self.text(r, c) == "" for c in range(self.ncols))


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def read_grid(path: str | Path) -> Grid:
    """Read the first sheet of a .xls, .xlsx or .csv file."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return _read_xls(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".csv":
        return _read_csv(path)
    raise UnsupportedFile(suffix)


def _read_xls(path: Path) -> Grid:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True)
    sheet = book.sheet_by_index(0)
    cells: list[list[object]] = []
    for r in range(sheet.nrows):
        row: list[object] = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                row.append(None)
            elif cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype == xlrd.XL_CELL_TEXT:
                row.append(cell.value if cell.value.strip() else None)
            else:
                row.append(cell.value)
        cells.append(row)
    return Grid(cells=cells, merged=[tuple(m) for m in sheet.merged_cells], source_format="xls")


def _read_xlsx(path: Path) -> Grid:
    import openpyxl

    book = openpyxl.load_workbook(str(path), read_only=False, data_only=True)
    sheet = book.worksheets[0]
    cells: list[list[object]] = []
    for row in sheet.iter_rows(values_only=True):
        cells.append([None if (isinstance(v, str) and not v.strip()) else v for v in row])
    merged = [
        (rng.min_row - 1, rng.max_row, rng.min_col - 1, rng.max_col)
        for rng in sheet.merged_cells.ranges
    ]
    book.close()
    return Grid(cells=cells, merged=merged, source_format="xlsx")


CSV_DELIMITERS = (",", ";", "\t", "|")


def guess_delimiter(text: str) -> str:
    """The delimiter that splits the header row and stays consistent over the next rows.

    Predictable on every Python version (csv.Sniffer's guesses differ between 3.12 and
    3.13). Indonesian Excel often saves CSV with ";" because "," is the decimal mark.
    """
    lines = [line for line in text.splitlines()[:50] if line.strip()]
    if not lines:
        return ","
    best, best_score = ",", (0.0, 0)
    for delim in CSV_DELIMITERS:
        fields = len(next(csv.reader([lines[0]], delimiter=delim)))
        if fields < 2:
            continue
        rows = [len(r) for r in csv.reader(lines[1:21], delimiter=delim)]
        consistent = sum(n == fields for n in rows) / len(rows) if rows else 1.0
        score = (consistent, fields)
        if score > best_score:
            best, best_score = delim, score
    return best


def _read_csv(path: Path) -> Grid:
    raw = path.read_bytes()
    text = _decode(raw)
    rows = list(csv.reader(io.StringIO(text), delimiter=guess_delimiter(text)))
    cells = [[v if v.strip() else None for v in row] for row in rows]
    return Grid(cells=cells, merged=[], source_format="csv")


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")
